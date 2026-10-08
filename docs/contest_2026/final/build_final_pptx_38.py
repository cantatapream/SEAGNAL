"""38장 배경(NotebookLM 그림) + 실제 캡처 덧붙이기 → 미리보기 PNG + PPTX."""
import os, sys
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.util import Emu

D = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(D, '..', 'images') + '/'  # docs/contest_2026/images
OUT = D + '/out/'
os.makedirs(OUT + 'ov', exist_ok=True)
os.makedirs(OUT + 'prev', exist_ok=True)
W, H = 1376, 768
K = 2  # 덧붙일 그림 해상도 배율

# mode: cover(꽉 채우고 넘치는 부분 자름, anchor top/center) | contain(전체가 보이게, 남는 곳 배경색)
# r: 모서리 둥글기(px, 1376 기준)  fade: 아래쪽 흐려짐 높이(px)
def P(box, f, mode='cover', anchor='top', r=14, fade=0, bg=None, pad=0, sy=None):
    # sy: cover 일 때 원본 그림에서 이 y(px)부터 보이게 (보여 줄 부분 지정)
    return dict(box=box, f=f, mode=mode, anchor=anchor, r=r, fade=fade, bg=bg, pad=pad, sy=sy)

L = {
 1: [P((632, 52, 744, 165), 'logo_SEAGNAL_앱아이콘.png', r=12)],
 2: [P((90, 74, 380, 694), 'S09-1_특보정보탭.png', r=34),
     P((944, 652, 1020, 728), 'S02-1_수상_해양경찰청AI경진대회.png', anchor='center', r=8)],
 4: [P((137, 234, 353, 290), 'logo_기상청.png', 'contain', r=8, bg=(255, 255, 255), pad=6),
     P((75, 310, 416, 372), 'logo_국립해양조사원.png', 'contain', r=8, bg=(255, 255, 255), pad=6),
     P((137, 388, 353, 450), 'logo_국가법령정보센터.png', 'contain', r=8, bg=(255, 255, 255), pad=8),
     P((1213, 552, 1268, 608), 'logo_윈디_모자이크.png', 'contain', r=8, bg=(255, 255, 255)),
     P((1278, 552, 1333, 608), 'logo_바다타임_모자이크.png', 'contain', r=8, bg=(255, 255, 255))],
 7: [P((1186, 50, 1307, 170), 'S02-1_수상_해양경찰청AI경진대회.png', anchor='center', r=10)],
 8: [P((714, 492, 1310, 720), '@S08', 'contain', r=10)],
 9: [P((127, 141, 237, 355), 'S09-1_특보정보탭.png', r=14),
     P((381, 141, 491, 355), 'S09-2_해양종합정보탭.png', r=14),
     P((635, 141, 745, 355), 'S09-3_해양안전_사고분석.png', r=14),
     P((889, 141, 999, 355), 'S09-4_해양생활_바다낚시.png', r=14),
     P((1145, 141, 1255, 355), 'S09-5_공지사항_전체.png', r=14)],
 10: [P((153, 160, 441, 704), 'S09-1_특보정보탭.png', r=34)],
 11: [P((214, 240, 498, 456), 'S11-1_해상기상전망.png', r=24, fade=40),
      P((884, 240, 1168, 392), 'S11-2_해상일기도.png', r=24, fade=30)],
 12: [P((573, 163, 805, 627), 'S12-2_특보구역카드_전체화면.png', 'contain', r=26, bg=(15, 23, 42))],
 13: [P((203, 160, 451, 500), 'S13-1_특보알림_크롭.png', r=24, fade=50),
      P((927, 160, 1175, 432), 'S13-3_안전권고팝업_발표.png', 'contain', anchor='top', r=24)],
 14: [P((158, 236, 376, 632), 'S14-1_챗봇답변_크롭.png', 'contain', r=24)],
 17: [P((124, 150, 354, 604), 'S09-2_해양종합정보탭.png', r=28)],
 18: [P((318, 144, 436, 368), 'S09-2_해양종합정보탭.png', r=14),
      P((630, 144, 746, 368), 'S18-2_유향유속.png', r=14),
      P((940, 144, 1058, 368), 'S18-3_파고파향.png', r=14)],
 20: [P((228, 212, 628, 406), 'S20-1_태풍_기상청.png', r=30, fade=24, sy=525),
      P((750, 212, 1147, 406), 'S20-2_태풍_JMA.png', r=30, fade=24, sy=556)],
 21: [P((110, 160, 420, 598), 'S09-3_해양안전_사고분석.png', r=30)],
 22: [P((408, 214, 598, 574), 'S22-1_위험지형_잠김경고.png', r=22),
      P((778, 214, 964, 574), 'S22-3_물빠짐_예측팝업.png', r=22)],
 24: [P((112, 222, 272, 536), 'S24-1_금지구역_지도.png', anchor='center', r=10),
      P((276, 222, 438, 536), 'S24-3_출입통제구역_팝업.png', r=10),
      P((531, 308, 758, 570), 'S24-4_항행경보.png', anchor='center', r=10),
      P((1088, 250, 1247, 361), 'S24-9_항로해역.png', anchor='center', r=8)],
 25: [P((250, 199, 298, 275), 'S25-2_바다낚시_팝업.png', 'contain', r=5),
     P((525, 199, 573, 275), 'S25-4_서핑_팝업.png', 'contain', r=5),
     P((801, 199, 850, 275), 'S25-10_바다갈라짐_7일표.png', 'contain', r=5),
     P((1077, 199, 1126, 275), 'S25-11_너울_해안선.png', r=5)],
 26: [P((250, 199, 298, 276), 'S09-5_공지사항_전체.png', r=5),
     P((662, 199, 711, 276), 'S26-1_제보_관리자답변.png', 'contain', r=5),
     P((1077, 199, 1126, 276), 'S26-2_설정_관심해역.png', 'contain', r=5)],
 30: [P((551, 427, 668, 546), 'logo_SEAGNAL_앱아이콘.png', r=12),
      P((708, 427, 825, 546), 'S30-1_앱다운로드_QR.png', 'contain', r=8, bg=(255, 255, 255), pad=4)],
}


# 슬라이드 위에 올리는 고칠 수 있는 글 상자: (x0, y0, x1, y1), 글, 크기(px, 1376 기준), 색
TXT = {}
FONT = '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc'

# 그림 속 문구를 고칠 때: 해당 자리를 바탕색 상자로 덮고 새 글을 올린다 (파워포인트에서 고칠 수 있음)
# (x0, y0, x1, y1), 바탕색, 글, 크기(px), 글자색
PATCH = {
}


def load(f):
    if f == '@S08':  # 댓글 두 장을 위아래로 쌓기
        s = Image.open(IMG + 'S08_사용자반응.png').convert('RGB')
        a, b = s.crop((0, 0, 845, 222)), s.crop((860, 0, 1705, 222))
        c = Image.new('RGB', (845, 454), s.getpixel((5, 5)))
        c.paste(a, (0, 0)); c.paste(b, (0, 232))
        return c
    return Image.open(IMG + f).convert('RGBA')


def bgcol(slide, box):
    x0, y0, x1, y1 = box
    c = slide.crop((x0 + (x1 - x0) // 3, y0 + (y1 - y0) // 3, x1 - (x1 - x0) // 3, y1 - (y1 - y0) // 3)).convert('RGB')
    return c.resize((1, 1), Image.BOX).getpixel((0, 0))


def render(slide, p):
    x0, y0, x1, y1 = p['box']
    bw, bh = (x1 - x0) * K, (y1 - y0) * K
    src = load(p['f'])
    if p['mode'] == 'cover':
        s = max(bw / src.width, bh / src.height)
        im = src.resize((round(src.width * s), round(src.height * s)), Image.LANCZOS)
        ox = (im.width - bw) // 2
        oy = 0 if p['anchor'] == 'top' else (im.height - bh) // 2
        if p['sy'] is not None:
            oy = max(0, min(im.height - bh, round(p['sy'] * s)))
        im = im.crop((ox, oy, ox + bw, oy + bh))
        canvas = Image.new('RGBA', (bw, bh), (0, 0, 0, 255))
        canvas.alpha_composite(im.convert('RGBA'))
    else:
        bg = p['bg'] or bgcol(slide, p['box'])
        canvas = Image.new('RGBA', (bw, bh), bg + (255,))
        pad = p['pad'] * K
        s = min((bw - 2 * pad) / src.width, (bh - 2 * pad) / src.height)
        im = src.resize((round(src.width * s), round(src.height * s)), Image.LANCZOS).convert('RGBA')
        ox = (bw - im.width) // 2
        oy = pad if p['anchor'] == 'top' else (bh - im.height) // 2
        canvas.alpha_composite(im, (ox, oy))
    m = Image.new('L', (bw, bh), 0)
    ImageDraw.Draw(m).rounded_rectangle((0, 0, bw - 1, bh - 1), radius=p['r'] * K, fill=255)
    if p['fade']:
        fh = p['fade'] * K
        g = Image.new('L', (bw, bh), 255)
        gd = ImageDraw.Draw(g)
        for i in range(fh):
            gd.line((0, bh - fh + i, bw, bh - fh + i), fill=int(255 * (1 - (i + 1) / fh)))
        m = Image.composite(m, Image.new('L', (bw, bh), 0), g) if False else Image.fromarray(
            __import__('numpy').minimum(__import__('numpy').asarray(m), __import__('numpy').asarray(g)))
    canvas.putalpha(m)
    return canvas


prs = Presentation()
prs.slide_width = Emu(16256000)
prs.slide_height = Emu(9144000)  # NotebookLM PPTX와 같은 16:9 크기
E = 16256000 / W
Y0, EY = 38100, 9067800 / H  # NotebookLM과 같게: 배경을 위아래 여백 38100 두고 비율 유지
blank = prs.slide_layouts[6]
SC = os.path.dirname(D)  # ⚠원본은 작업 폴더(scratchpad)의 NotebookLM 추출 그림 r1cpx·r2dpx·r2bpx·px — 저장소에는 없음
FIX = {2: 'r1cpx/s1_image1.png', 7: 'r1cpx/s2_image2.png', 15: 'r1cpx/s3_image3.png', 16: 'r1cpx/s4_image4.png',
       28: 'r1cpx/s5_image5.png', 29: 'r1epx/s6.png', 30: 'r1cpx/s7_image7.png'}
NEW = {'A1': 'r2dpx/s1.png', 'A2': 'r2dpx/s2.png', 'A3': 'r2dpx/s3.png', 'A4': 'r2dpx/s4.png',
       'A5': 'r2dpx/s5.png', 'A6': 'r2bpx/s6.png', 'A7': 'r2dpx/s7.png', 'A8': 'r2dpx/s8.png'}
ORDER = [1, 2, 'A1', 3, 4, 5, 6, 7, 8, 'A2', 'A3', 'A4', 'A5', 9, 10, 11, 12, 13, 'A8'] + list(range(14, 27)) + ['A6', 27, 28, 'A7', 29, 30]
# 다시 만든 장은 빈 틀 위치가 바뀌어 새로 잰 값으로 바꾼다
L[2] = [P((101, 79, 396, 690), 'S09-1_특보정보탭.png', r=32),
        P((919, 607, 1006, 694), 'S02-1_수상_해양경찰청AI경진대회.png', anchor='center', r=6)]
L[7] = [P((1240, 53, 1318, 132), 'S02-1_수상_해양경찰청AI경진대회.png', anchor='center', r=10)]
L[30] = [P((560, 444, 675, 559), 'logo_SEAGNAL_앱아이콘.png', r=12),
         P((700, 444, 815, 559), 'S30-1_앱다운로드_QR.png', 'contain', r=8, bg=(255, 255, 255), pad=4)]
L['A2'] = [P((838, 376, 1304, 534), 'S_AI_제미나이_아이콘모음.png', 'contain', r=8)]
LINES = {28: [(5106196, 3050383, 0, 36000, 31750, (0xEE, 0xF1, 0xFA))]}  # 사용자가 PPT에서 넣은 「통」 획
for k, n in enumerate(ORDER, 1):
    src = NEW[n] if isinstance(n, str) else FIX.get(n)
    slide_img = Image.open(f'{SC}/{src}' if src else f'{D}/px/s{n:02d}.png').convert('RGB')
    sl = prs.slides.add_slide(blank)
    bgj = f'{OUT}ov/bg{k:02d}.jpg'; slide_img.save(bgj, quality=92)
    sl.shapes.add_picture(bgj, 0, Emu(Y0), prs.slide_width, Emu(9067800))
    prev = slide_img.convert('RGBA')
    for i, p in enumerate(L.get(n, [])):
        ov = render(slide_img, p)
        fn = f'{OUT}ov/s{k:02d}_{i}.png'
        ov.save(fn)
        x0, y0, x1, y1 = p['box']
        pic = sl.shapes.add_picture(fn, Emu(round(x0 * E)), Emu(round(Y0 + y0 * EY)), Emu(round((x1 - x0) * E)), Emu(round((y1 - y0) * EY)))
        pic.name = os.path.splitext(p['f'].lstrip('@'))[0]
        prev.alpha_composite(ov.resize((x1 - x0, y1 - y0), Image.LANCZOS), (x0, y0))
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from PIL import ImageFont
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    for lx, ly, lcx, lcy, lw, lc in LINES.get(n, []):
        from pptx.enum.shapes import MSO_CONNECTOR
        ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(lx), Emu(ly), Emu(lx + lcx), Emu(ly + lcy))
        ln.line.width = Emu(lw); ln.line.color.rgb = RGBColor(*lc)
        px0 = lx / E; py0 = (ly - Y0) / EY
        ImageDraw.Draw(prev).line([(px0, py0), (px0, py0 + lcy / EY)], fill=lc + (255,), width=max(1, round(lw / E)))
    for (x0, y0, x1, y1), bgc, t, px, col in PATCH.get(n, []):
        sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(round(x0 * E)), Emu(round(Y0 + y0 * EY)), Emu(round((x1 - x0) * E)), Emu(round((y1 - y0) * EY)))
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(*bgc); sh.line.fill.background(); sh.shadow.inherit = False
        tf = sh.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = tf.paragraphs[0]; para.alignment = 1
        r = para.add_run(); r.text = t; r.font.size = Pt(px * 1280 / W); r.font.bold = True
        r.font.name = '맑은 고딕'; r.font.color.rgb = RGBColor(*col)
        dd = ImageDraw.Draw(prev); dd.rectangle((x0, y0, x1, y1), fill=bgc + (255,))
        fnt = ImageFont.truetype(FONT, px); dd.text((x0, (y0 + y1) // 2), t, font=fnt, fill=col + (255,), anchor='lm')
    for (x0, y0, x1, y1), t, px, col in TXT.get(n, []):
        tb = sl.shapes.add_textbox(Emu(round(x0 * E)), Emu(round(Y0 + y0 * EY)), Emu(round((x1 - x0) * E)), Emu(round((y1 - y0) * EY)))
        tf = tb.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.word_wrap = False
        r = tf.paragraphs[0].add_run(); r.text = t
        r.font.size = Pt(px * 1280 / W); r.font.name = '맑은 고딕'; r.font.color.rgb = RGBColor(*col)
        ImageDraw.Draw(prev).text((x0, y0), t, font=ImageFont.truetype(FONT, px), fill=col + (255,))
    prev.convert('RGB').save(f'{OUT}prev/s{k:02d}.jpg', quality=88)
prs.core_properties.title = '바다 : 그 날의 신호, SEA:GNAL'
prs.save(OUT + 'tmp.pptx')
# 템플릿에 남은 4:3 표기·빈 미리보기 그림 정리
import zipfile, re, io
zin = zipfile.ZipFile(OUT + 'tmp.pptx')
zout = zipfile.ZipFile(OUT + 'SEAGNAL_final_38.pptx', 'w', zipfile.ZIP_DEFLATED)
th = io.BytesIO(); Image.open(f'{OUT}prev/s01.jpg').resize((256, 143)).save(th, 'JPEG', quality=85)
for it in zin.infolist():
    data = zin.read(it.filename)
    if it.filename == 'ppt/presentation.xml':
        data = re.sub(rb'<p:sldSz [^>]*/>', b'<p:sldSz cx="16256000" cy="9144000"/>', data)
    elif it.filename == 'docProps/app.xml':
        data = data.replace(b'On-screen Show (4:3)', b'Custom').replace(b'<Slides>0</Slides>', b'<Slides>38</Slides>')
    elif it.filename == 'docProps/thumbnail.jpeg':
        data = th.getvalue()
    zout.writestr(it, data)
zout.close(); os.remove(OUT + 'tmp.pptx')
print('ok', os.path.getsize(OUT + 'SEAGNAL_final_38.pptx'))
